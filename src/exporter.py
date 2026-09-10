"""Export Pipeline Datasets to Multi-Tab Excel Workbook and Standalone CSVs."""

import csv
import logging
import os
from typing import List
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.models.schemas import (
    StartupEntity,
    ProductEntity,
    ResearchPaperEntity,
    JobEntity,
    NewsEntity,
    EntityMappingLog
)

logger = logging.getLogger("DatasetExporter")


class DatasetExporter:
    """Exports structured entities into multi-tab Excel workbook and individual CSV files."""

    def __init__(self, output_dir: str = "output"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)

    def export_all(
        self,
        startups: List[StartupEntity],
        products: List[ProductEntity],
        papers: List[ResearchPaperEntity],
        jobs: List[JobEntity],
        news: List[NewsEntity],
        mapping_logs: List[EntityMappingLog],
        excel_filename: str = "pipeline_output.xlsx"
    ) -> str:
        """Export all datasets to a single 6-tab Excel workbook and separate CSV files."""
        excel_path = os.path.join(self.output_dir, excel_filename)
        logger.info(f"Exporting pipeline data to {excel_path} and CSVs in {self.output_dir}...")

        # 1. Flatten into records lists
        startups_rows = [s.to_flat_dict() for s in startups]
        products_rows = [p.to_flat_dict() for p in products]
        papers_rows = [r.to_flat_dict() for r in papers]
        jobs_rows = [j.to_flat_dict() for j in jobs]
        news_rows = [n.to_flat_dict() for n in news]
        logs_rows = [m.to_flat_dict() for m in mapping_logs]

        # 2. Write CSVs
        self._write_csv("startups.csv", startups_rows)
        self._write_csv("products.csv", products_rows)
        self._write_csv("research_papers.csv", papers_rows)
        self._write_csv("jobs.csv", jobs_rows)
        self._write_csv("news.csv", news_rows)
        self._write_csv("entity_mapping_log.csv", logs_rows)

        # 3. Create Multi-tab Excel Workbook with formatting
        wb = Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        sheets_data = [
            ("Startups", startups_rows),
            ("Products", products_rows),
            ("Research Papers", papers_rows),
            ("Jobs", jobs_rows),
            ("News", news_rows),
            ("Entity Mapping Log", logs_rows)
        ]

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
        border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0")
        )

        for sheet_title, rows in sheets_data:
            ws = wb.create_sheet(title=sheet_title)
            ws.views.sheetView[0].showGridLines = True

            if not rows:
                ws.append(["No records collected"])
                continue

            headers = list(rows[0].keys())
            ws.append(headers)

            # Style header row
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Append data rows
            for r in rows:
                row_values = [str(r.get(h, "")) if r.get(h) is not None else "" for h in headers]
                ws.append(row_values)

            # Auto-fit column widths
            for col_idx, col_cells in enumerate(ws.columns, 1):
                max_len = max(len(str(cell.value or "")) for cell in col_cells[:100])
                col_letter = get_column_letter(col_idx)
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

        wb.save(excel_path)
        logger.info(f"Excel workbook successfully saved to {excel_path}")
        return excel_path

    def _write_csv(self, filename: str, rows: List[dict]):
        """Write records to a CSV file."""
        if not rows:
            return
        filepath = os.path.join(self.output_dir, filename)
        keys = list(rows[0].keys())
        with open(filepath, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=keys)
            writer.writeheader()
            writer.writerows(rows)
        logger.info(f"Wrote {len(rows)} rows to {filepath}")
